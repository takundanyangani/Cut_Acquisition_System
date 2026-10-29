"""Excel workbooks. CUT ICT 38 §5.7 records requests in Excel for tracking;
the export keeps that working without anyone retyping data."""

from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.core.workdays import working_days_between

from . import workflow
from .models import State

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEADER_FILL = PatternFill("solid", fgColor="E7E2D6")


def _local(value):
    if value is None:
        return None
    return timezone.localtime(value).replace(tzinfo=None)


def _sheet(workbook, title, headers, widths):
    sheet = workbook.active if workbook.active.title == "Sheet" else workbook.create_sheet()
    sheet.title = title
    sheet.append(headers)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
        cell = sheet.cell(row=1, column=index)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    return sheet


def _finish(sheet, date_columns, datetime_columns):
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if cell.column in datetime_columns and cell.value:
                cell.number_format = "dd/mm/yyyy hh:mm"
            elif cell.column in date_columns and cell.value:
                cell.number_format = "dd/mm/yyyy"
    sheet.auto_filter.ref = sheet.dimensions


def _response(workbook, name):
    buffer = BytesIO()
    workbook.save(buffer)
    stamp = timezone.localtime().strftime("%Y%m%d-%H%M")
    response = HttpResponse(buffer.getvalue(), content_type=XLSX)
    response["Content-Disposition"] = f'attachment; filename="{name}-{stamp}.xlsx"'
    response["Cache-Control"] = "private, no-store"
    return response


def requests_workbook(queryset):
    workbook = Workbook()
    headers = [
        "Reference",
        "Procedure",
        "Type",
        "Title",
        "Items",
        "Quantity",
        "Requester",
        "Department",
        "Location",
        "Priority",
        "Stage",
        "With",
        "Submitted",
        "Director approved",
        "Sent to Procurement",
        "Procurement ref",
        "Delivery complete",
        "Closed",
        "Working days open",
        "Overdue",
    ]
    widths = [16, 12, 22, 34, 40, 9, 24, 24, 30, 10, 28, 26, 17, 17, 17, 16, 17, 17, 10, 9]
    sheet = _sheet(workbook, "Requests", headers, widths)
    now = timezone.now()
    for req in queryset.select_related("requester", "department", "campus", "assigned_technician"):
        end = req.closed_at or now
        sheet.append(
            [
                req.reference,
                req.procedure_number,
                req.get_kind_display(),
                req.title,
                req.items_summary,
                req.total_quantity,
                req.requester.display_name,
                str(req.department) if req.department_id else "",
                req.location_text,
                req.get_priority_display(),
                State(req.state).label,
                workflow.stage_holder(req),
                _local(req.submitted_at),
                _local(req.approved_at),
                _local(req.sent_to_procurement_at),
                req.procurement_reference,
                _local(req.delivery_completed_at),
                _local(req.closed_at),
                working_days_between(req.submitted_at, end) if req.submitted_at else None,
                "Yes" if req.is_overdue else "",
            ]
        )
    _finish(sheet, set(), {13, 14, 15, 17, 18})
    return _response(workbook, "ict-requests")


def assets_workbook(queryset):
    workbook = Workbook()
    headers = [
        "Asset tag",
        "Serial number",
        "Type",
        "Make",
        "Model",
        "Specification",
        "Condition",
        "Status",
        "Holder",
        "Location",
        "Supplier",
        "Purchase order",
        "Acquired",
        "Warranty ends",
        "Source request",
    ]
    widths = [16, 20, 16, 14, 22, 40, 10, 24, 24, 30, 20, 16, 12, 12, 16]
    sheet = _sheet(workbook, "Assets", headers, widths)
    for asset in queryset.select_related("category", "holder", "campus", "source_requisition"):
        sheet.append(
            [
                asset.asset_tag,
                asset.serial_number or "",
                asset.category.name,
                asset.make,
                asset.model,
                asset.specification,
                asset.get_condition_display(),
                asset.get_status_display(),
                asset.holder.display_name if asset.holder_id else "",
                asset.location_text,
                asset.supplier,
                asset.purchase_order_ref,
                asset.acquired_on,
                asset.warranty_expires_on,
                asset.source_requisition.reference if asset.source_requisition_id else "",
            ]
        )
    _finish(sheet, {13, 14}, set())
    return _response(workbook, "ict-asset-register")
